from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PRD.md"
OUTPUT = ROOT / "激光训练屋产品需求文档-V1.0.docx"

INK = "17201B"
GREEN = "B7DD43"
DARK = "111719"
MID = "49544F"
LIGHT = "EEF2EC"
MUTED = "68716C"
WHITE = "FFFFFF"


def font(run, name="STHeitiSC-Medium", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_inline(paragraph, text, color=INK):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            font(run, "Consolas", 9, False, "39413D")
            run.font.highlight_color = None
        elif part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            font(run, size=10.5, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            font(run, size=10.5, color=color)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "STHeitiSC-Medium"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeitiSC-Medium")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    for name, size, color, before, after in (
        ("Heading 1", 17, DARK, 15, 6),
        ("Heading 2", 13.5, DARK, 11, 4),
        ("Heading 3", 11.5, MID, 8, 3),
    ):
        st = doc.styles[name]
        st.font.name = "STHeitiSC-Medium"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeitiSC-Medium")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(62)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("NBYS  TRAINING SYSTEM")
    font(r, "Arial", 11, True, GREEN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("激光训练屋")
    font(r, size=30, bold=True, color=DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("产品需求文档  /  PRODUCT REQUIREMENTS")
    font(r, "Arial", 13, True, MID)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Inches(1.55), Inches(4.8)]
    values = [
        ("文档版本", "V1.0"),
        ("产品阶段", "第一阶段：产品需求评审"),
        ("适用版本", "训练屋 MVP"),
        ("目标终端", "H5 会员端 / Android 靶机 / Web 后管"),
        ("评审状态", "待确认，未进入 UX 与开发"),
    ]
    for row, pair in zip(table.rows, values):
        for idx, value in enumerate(pair):
            row.cells[idx].width = widths[idx]
            margins(row.cells[idx], 130, 150, 130, 150)
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(row.cells[idx], DARK if idx == 0 else LIGHT)
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            font(run, size=9.5 if idx == 0 else 10.5, bold=idx == 0, color=WHITE if idx == 0 else INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("评审门禁")
    font(r, size=11, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, "本文件确认后才进入 UX 原型；三端原型再次确认后才允许开发正式代码。", DARK)

    doc.add_page_break()


def add_table(doc, rows):
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = 6.62
    if columns == 2:
        widths = [total * 0.28, total * 0.72]
    elif columns == 3:
        widths = [total * 0.22, total * 0.34, total * 0.44]
    else:
        widths = [total / columns] * columns
    for ridx, values in enumerate(rows):
        row = table.add_row()
        if ridx == 0:
            repeat_header(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            shade(cell, DARK if ridx == 0 else ("F5F7F3" if ridx % 2 == 0 else WHITE))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            run = p.add_run(value)
            font(run, size=8.6 if columns >= 4 else 9, bold=ridx == 0, color=WHITE if ridx == 0 else INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def build():
    doc = Document()
    style_document(doc)
    add_cover(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_mermaid = False
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    for raw in lines:
        line = raw.strip()
        if line.startswith("# 激光训练屋产品需求文档"):
            continue
        if line.startswith("```mermaid"):
            flush_table()
            in_mermaid = True
            continue
        if in_mermaid:
            if line == "```":
                in_mermaid = False
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            table_rows.append(cells)
            continue
        flush_table()
        if not line:
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:], DARK)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], DARK)
        elif line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:], MID)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
    flush_table()

    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("NBYS  /  激光训练屋 PRD  V1.0")
    font(r, "Arial", 8, True, MUTED)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("产品评审稿 · 原型确认前禁止进入正式开发")
    font(r, size=8, color=MUTED)

    doc.core_properties.title = "激光训练屋产品需求文档 V1.0"
    doc.core_properties.subject = "H5、Android 靶机与 Web 后管的训练屋 MVP 产品需求"
    doc.core_properties.author = "NBYS 产品团队"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
